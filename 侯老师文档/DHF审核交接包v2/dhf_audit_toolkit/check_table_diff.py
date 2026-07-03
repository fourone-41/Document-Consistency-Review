# -*- coding: utf-8 -*-
"""
Tier1 形式化算法 · 行级 key-join 精确 diff（确定性，无需 AI）
================================================================
比两份同构表(BOM/组件清单/版本前后)：按主键(物料编码/图号)对齐，输出
  仅A有 / 仅B有 / 两边都有但某列值不同(逐列给 A值 vs B值)。
治"同一份BOM在两阶段/两版本之间是否漂移"——给可签字的定位级差异。

用法:
  python check_table_diff.py A.xls B.xls --key 物料编码 --cols 材质,规格,型号,封装形式,配比 --out diff.csv
  (--key/--cols/--header 不给则自动探测)
"""
import sys, os, csv, argparse
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass

def read_table(path):
    ext=os.path.splitext(path)[1].lower()
    rows=[]
    if ext==".xlsx":
        import openpyxl
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        sh=wb[wb.sheetnames[0]]
        for r in sh.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v).strip() for v in r])
        wb.close()
    else:
        import xlrd
        wb=xlrd.open_workbook(path); sh=wb.sheet_by_index(0)
        for ri in range(sh.nrows):
            rows.append([str(sh.cell_value(ri,ci)).strip() for ci in range(sh.ncols)])
    return rows

def find_header(rows):
    for i,r in enumerate(rows[:15]):
        joined="".join(r)
        if ("物料编码" in joined or "物料编号" in joined) and ("序号" in joined or "品" in joined or "图号" in joined):
            return i
        if "图号" in joined and ("名称" in joined or "品" in joined):
            return i
    return 0

def norm(s): return " ".join((s or "").split()).replace("　"," ").strip()

def load(path, key_hint, header_hint):
    rows=read_table(path)
    h = header_hint if header_hint is not None else find_header(rows)
    cols=[norm(c) for c in rows[h]]
    # 主键列
    keycands=[key_hint] if key_hint else ["物料编码","物料编号","图号","料号"]
    kidx=None
    for kc in keycands:
        for ci,c in enumerate(cols):
            if kc and kc in c: kidx=ci; break
        if kidx is not None: break
    recs={}; dup=[]
    for r in rows[h+1:]:
        if len(r)<len(cols): r=r+[""]*(len(cols)-len(r))
        d={cols[ci]:norm(r[ci]) for ci in range(len(cols))}
        key = norm(r[kidx]) if kidx is not None and kidx<len(r) else ""
        if not key:  # 主键空:用 品名+型号 兜底;仍空=章节/空行,跳过
            nm=next((d[c] for c in cols if "品" in c or "名称" in c),"")
            tp=next((d[c] for c in cols if "型号" in c),"")
            key=(nm+"|"+tp).strip("|")
        if not key or key in ("nan",): continue
        if key in recs: dup.append(key)
        else: recs[key]=d
    return cols, kidx, recs, dup

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("a"); ap.add_argument("b")
    ap.add_argument("--key",default=""); ap.add_argument("--cols",default="")
    ap.add_argument("--header",type=int,default=None)
    ap.add_argument("--labelA",default="A"); ap.add_argument("--labelB",default="B")
    ap.add_argument("--out",default="表diff.csv")
    a=ap.parse_args()

    colsA,kA,recA,dupA=load(a.a,a.key,a.header)
    colsB,kB,recB,dupB=load(a.b,a.key,a.header)
    keyname = (colsA[kA] if kA is not None else a.key) or "key"
    # 要比的列:交集 - 主键 - 序号
    cmp_cols=[c for c in a.cols.split(",") if c.strip()] if a.cols else \
             [c for c in colsA if c in colsB and c!=keyname and "序号" not in c and c!=""]
    onlyA=sorted(set(recA)-set(recB)); onlyB=sorted(set(recB)-set(recA)); both=sorted(set(recA)&set(recB))
    diffs=[]
    for k in both:
        for c in cmp_cols:
            va=recA[k].get(c,""); vb=recB[k].get(c,"")
            if norm(va)!=norm(vb) and (va or vb):
                diffs.append({"主键":k,"列":c,a.labelA:va,a.labelB:vb})
    # 输出
    with open(a.out,"w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f); w.writerow(["类型","主键("+keyname+")","列",a.labelA,a.labelB])
        for k in onlyA: w.writerow(["仅"+a.labelA+"有",k,"",recA[k].get(next((c for c in colsA if '品' in c or '名称' in c),''),''),""])
        for k in onlyB: w.writerow(["仅"+a.labelB+"有",k,"","",recB[k].get(next((c for c in colsB if '品' in c or '名称' in c),''),'')])
        for d in diffs: w.writerow(["值不同",d["主键"],d["列"],d[a.labelA],d[a.labelB]])
    print(f"A={a.labelA}: {os.path.basename(a.a)}  ({len(recA)}行, 重复键{len(set(dupA))})")
    print(f"B={a.labelB}: {os.path.basename(a.b)}  ({len(recB)}行, 重复键{len(set(dupB))})")
    print(f"主键={keyname} | 比对列={cmp_cols}")
    print(f"  仅{a.labelA}有: {len(onlyA)}  仅{a.labelB}有: {len(onlyB)}  值不同: {len(diffs)} 处")
    if onlyA: print("  仅"+a.labelA+"有:", "; ".join(onlyA[:12]))
    if onlyB: print("  仅"+a.labelB+"有:", "; ".join(onlyB[:12]))
    for d in diffs[:25]: print(f"  ▲ [{d['列']}] {d['主键']}: {a.labelA}={d[a.labelA]!r} vs {a.labelB}={d[a.labelB]!r}")
    print("  ->",a.out)

if __name__=="__main__":
    main()
