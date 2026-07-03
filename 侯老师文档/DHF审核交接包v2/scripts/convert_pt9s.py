# -*- coding: utf-8 -*-
"""
PT9S 工具层：把 DHF 原始件(.doc/.docx/.xls/.xlsx/.pdf) 转成 Markdown，镜像目录结构。
- .doc/.docx : Word COM (复用单个 Word 实例，逐个开关) -> 正文文本 + 勾选框/OLE 计数(损失标记)
- .xls       : xlrd  -> 每个 sheet 转 MD 表
- .xlsx      : openpyxl -> 每个 sheet 转 MD 表
- .pdf       : fitz -> 每页文本；正文极少的页(疑似扫描/图/勾选框)标记 low_text 并渲染 PNG
每个 .md 带来源头部 + 损失标记；同时出一张总转换台账 _manifest_conversion.csv。
"""
import sys, os, csv, glob, json, re, datetime as dt
sys.stdout.reconfigure(encoding="utf-8")

SRC = r"D:\AI_0415\03DHF\PT9S\PT9S_DHF"
DST = r"D:\AI_0415\03DHF\PT9S_MD"
MANIFEST = os.path.join(DST, "_manifest_conversion.csv")
LOWTEXT_PDF_THRESHOLD = 40   # 一页正文少于此字符 -> 渲染PNG并标 low_text

def now(): return dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")
def clean(t):
    if not t: return ""
    t = t.replace("\x07"," | ").replace("\x0b","\n").replace("\r","\n")
    t = re.sub(r"[\x00-\x08\x0e-\x1f]", "", t)
    t = re.sub(r"[ \t]+"," ", t)
    t = re.sub(r"\n{3,}","\n\n", t)
    return t.strip()

def out_paths(src_file):
    rel = os.path.relpath(src_file, SRC)
    md = os.path.join(DST, os.path.splitext(rel)[0] + ".md")
    assets = os.path.join(DST, os.path.splitext(rel)[0] + ".assets")
    return rel, md, assets

def header(rel, src_file, typ, status, loss):
    return ("---\n"
            f"source_file: {src_file}\n"
            f"source_relative: {rel}\n"
            f"source_type: {typ}\n"
            f"conversion_status: {status}\n"
            f"loss_flags: {';'.join(loss) if loss else ''}\n"
            f"converted_at: {now()}\n"
            "manual_review_required: " + ("yes" if loss else "no") + "\n"
            "---\n\n")

def md_table(rows, max_rows=400, max_cols=20):
    rows = [r[:max_cols] for r in rows[:max_rows]]
    if not rows: return ""
    w = max(len(r) for r in rows)
    rows = [list(r)+[""]*(w-len(r)) for r in rows]
    def cell(x): return clean(str(x)).replace("\n"," ").replace("|","\\|")[:200]
    out = ["| " + " | ".join(cell(c) for c in rows[0]) + " |",
           "| " + " | ".join("---" for _ in rows[0]) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(cell(c) for c in r) + " |")
    note = ""
    if len(rows) >= max_rows: note = f"\n\n> (表过长,仅取前{max_rows}行)"
    return "\n".join(out) + note

# ---------------- Excel ----------------
def conv_xls(f):
    import xlrd
    wb = xlrd.open_workbook(f)
    parts, loss = [], []
    for name in wb.sheet_names():
        sh = wb.sheet_by_name(name)
        if sh.nrows == 0: continue
        rows = [[sh.cell_value(r,c) for c in range(sh.ncols)] for r in range(sh.nrows)]
        parts.append(f"### sheet: {name}  ({sh.nrows}x{sh.ncols})\n\n" + md_table(rows))
    return "\n\n".join(parts), loss

def conv_xlsx(f):
    import openpyxl
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    parts, loss = [], []
    for name in wb.sheetnames:
        sh = wb[name]
        rows = [[("" if v is None else v) for v in row] for row in sh.iter_rows(values_only=True)]
        rows = [r for r in rows if any(str(x).strip() for x in r)]
        if not rows: continue
        parts.append(f"### sheet: {name}  ({len(rows)}行)\n\n" + md_table(rows))
    wb.close()
    return "\n\n".join(parts), loss

# ---------------- PDF ----------------
def conv_pdf(f, assets):
    import fitz
    d = fitz.open(f); parts, loss = [], []
    low_pages = 0
    for i in range(d.page_count):
        pg = d[i]; t = pg.get_text().strip()
        if len(t) < LOWTEXT_PDF_THRESHOLD:
            low_pages += 1
            os.makedirs(assets, exist_ok=True)
            png = os.path.join(assets, f"page_{i+1:03d}.png")
            try:
                pg.get_pixmap(dpi=150).save(png)
                parts.append(f"### page {i+1} (正文少,已渲染PNG)\n\n![p{i+1}]({os.path.relpath(png, os.path.dirname(assets))})\n\n```\n{t}\n```")
            except Exception as e:
                parts.append(f"### page {i+1} (正文少,PNG失败:{e})\n\n```\n{t}\n```")
        else:
            parts.append(f"### page {i+1}\n\n{clean(t)}")
    if low_pages: loss.append(f"pdf_low_text_pages={low_pages}")
    d.close()
    return "\n\n".join(parts), loss

# ---------------- Word (COM, 复用实例) ----------------
class WordConv:
    def __init__(self):
        import pythoncom, win32com.client
        self.pythoncom = pythoncom; pythoncom.CoInitialize()
        self.word = win32com.client.DispatchEx("Word.Application")
        self.word.Visible = False; self.word.DisplayAlerts = 0
    def conv(self, f):
        doc = self.word.Documents.Open(os.path.abspath(f), ReadOnly=True,
                                       AddToRecentFiles=False, ConfirmConversions=False)
        loss = []
        try:
            text = clean(str(doc.Range().Text or ""))
            ntab = int(doc.Tables.Count)
            # 勾选框
            cb=cbc=0
            try:
                for i in range(1, int(doc.FormFields.Count)+1):
                    ff=doc.FormFields.Item(i)
                    try:
                        v=ff.CheckBox; cb+=1
                        if bool(v.Value): cbc+=1
                    except Exception: pass
            except Exception: pass
            # OLE / 浮动形状
            ole=0
            try:
                for i in range(1,int(doc.InlineShapes.Count)+1):
                    sh=doc.InlineShapes.Item(i)
                    try:
                        if int(sh.Type) in (1,2,5): ole+=1
                    except Exception: pass
            except Exception: pass
            shapes=0
            try: shapes=int(doc.Shapes.Count)
            except Exception: pass
            if cb>0: loss.append(f"checkbox_form_fields={cb}(checked={cbc})")
            if ole>0: loss.append(f"embedded_ole={ole}")
            if shapes>0: loss.append(f"floating_shapes={shapes}")
            body = f"> 表格数:{ntab}  勾选框:{cb}(勾选{cbc})  OLE:{ole}  浮动形状:{shapes}\n\n{text}"
            return body, loss
        finally:
            doc.Close(False)
    def close(self):
        try: self.word.Quit()
        except Exception: pass
        try: self.pythoncom.CoUninitialize()
        except Exception: pass

def main():
    os.makedirs(DST, exist_ok=True)
    files = [p for p in glob.glob(os.path.join(SRC,"**","*"), recursive=True)
             if os.path.isfile(p) and not os.path.basename(p).startswith("~$")]
    EXT = {".doc",".docx",".xls",".xlsx",".pdf"}
    files = [f for f in files if os.path.splitext(f)[1].lower() in EXT]
    files.sort()
    print(f"待转 {len(files)} 个文件 -> {DST}", flush=True)

    wc = None
    rows = []
    for idx,f in enumerate(files,1):
        rel, md, assets = out_paths(f)
        ext = os.path.splitext(f)[1].lower()
        stage = rel.split(os.sep)[0]
        status="failed"; body=""; loss=[]; err=""
        try:
            if ext in (".doc",".docx"):
                if wc is None: wc = WordConv()
                body, loss = wc.conv(f)
            elif ext==".xls":
                body, loss = conv_xls(f)
            elif ext==".xlsx":
                body, loss = conv_xlsx(f)
            elif ext==".pdf":
                body, loss = conv_pdf(f, assets)
            status = "ok" if body and body.strip() else "empty"
        except Exception as e:
            err = repr(e)[:300]; loss.append("convert_error")
        if status=="ok":
            os.makedirs(os.path.dirname(md), exist_ok=True)
            with open(md,"w",encoding="utf-8") as fh:
                fh.write(header(rel,f,ext.lstrip("."),status,loss)+body)
        rows.append({"rel":rel,"stage":stage,"type":ext.lstrip("."),"status":status,
                     "text_len":len(body),"loss_flags":";".join(loss),"md":md if status=="ok" else "",
                     "error":err})
        if idx%10==0 or idx==len(files):
            print(f"  {idx}/{len(files)} {status} {rel[:60]}", flush=True)
    if wc: wc.close()

    with open(MANIFEST,"w",encoding="utf-8-sig",newline="") as fh:
        w=csv.DictWriter(fh, fieldnames=["rel","stage","type","status","text_len","loss_flags","md","error"])
        w.writeheader(); w.writerows(rows)
    from collections import Counter
    print("\n=== 转换完成 ===")
    print("状态:", dict(Counter(r["status"] for r in rows)))
    print("类型:", dict(Counter(r["type"] for r in rows)))
    fails=[r for r in rows if r["status"]!="ok"]
    print(f"失败/空 {len(fails)} 个:")
    for r in fails[:30]: print("  -", r["status"], r["rel"], r["error"][:80])
    print("台账:", MANIFEST)

if __name__=="__main__":
    raise SystemExit(main())
