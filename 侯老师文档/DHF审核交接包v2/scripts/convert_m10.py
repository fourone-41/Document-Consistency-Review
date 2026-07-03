# -*- coding: utf-8 -*-
"""M10 分片并行转换器。一个进程转一个分片(按文件索引取模),多进程并行。
用法: python convert_m10.py --types pdf,xls,xlsx --shard 0/4
      python convert_m10.py --types doc,docx --shard 0/2
Word(.doc/.docx)走COM,每个进程独立实例;PDF/Excel纯Python可多开。
"""
import sys, os, csv, re, glob, argparse, datetime as dt
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass

ROOT=r"D:\AI_0415\03DHF\雾化器m10"
DST =r"D:\AI_0415\03DHF\m10_MD"

def now(): return dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")
def clean(t):
    if not t: return ""
    t=t.replace("\x07"," | ").replace("\x0b","\n").replace("\r","\n")
    t=re.sub(r"[\x00-\x08\x0e-\x1f]","",t); t=re.sub(r"[ \t]+"," ",t); t=re.sub(r"\n{3,}","\n\n",t)
    return t.strip()
def md_table(rows,max_rows=400,max_cols=20):
    rows=[list(r)[:max_cols] for r in rows[:max_rows]]
    if not rows: return ""
    w=max(len(r) for r in rows); rows=[list(r)+[""]*(w-len(r)) for r in rows]
    def c(x): return clean(str(x)).replace("\n"," ").replace("|","\\|")[:200]
    out=["| "+" | ".join(c(x) for x in rows[0])+" |","| "+" | ".join("---" for _ in rows[0])+" |"]
    for r in rows[1:]: out.append("| "+" | ".join(c(x) for x in r)+" |")
    return "\n".join(out)
def header(rel,src,typ,st,loss):
    return ("---\n"+f"source_file: {src}\nsource_relative: {rel}\nsource_type: {typ}\n"
            f"conversion_status: {st}\nloss_flags: {';'.join(loss)}\nconverted_at: {now()}\n---\n\n")

_word=None
def get_word():
    global _word
    if _word is None:
        import pythoncom, win32com.client
        pythoncom.CoInitialize()
        _word=win32com.client.DispatchEx("Word.Application"); _word.Visible=False; _word.DisplayAlerts=0
    return _word

def convert(f, ext):
    body="";loss=[];st="failed";err=""
    try:
        if ext in (".doc",".docx"):
            w=get_word()
            doc=w.Documents.Open(os.path.abspath(f),ReadOnly=True,AddToRecentFiles=False,ConfirmConversions=False)
            try:
                text=clean(str(doc.Range().Text or "")); ntab=int(doc.Tables.Count); cb=cbc=0
                try:
                    for k in range(1,int(doc.FormFields.Count)+1):
                        try:
                            v=doc.FormFields.Item(k).CheckBox; cb+=1; cbc+=1 if bool(v.Value) else 0
                        except Exception: pass
                except Exception: pass
                if cb: loss.append(f"checkbox={cb}(checked={cbc})")
                body=f"> 表格数:{ntab} 勾选框:{cb}(勾{cbc})\n\n{text}"
            finally: doc.Close(False)
        elif ext==".xls":
            import xlrd; wb=xlrd.open_workbook(f); parts=[]
            for nm in wb.sheet_names():
                sh=wb.sheet_by_name(nm)
                if sh.nrows: parts.append(f"### sheet: {nm}\n\n"+md_table([[sh.cell_value(r,c) for c in range(sh.ncols)] for r in range(sh.nrows)]))
            body="\n\n".join(parts)
        elif ext==".xlsx":
            import openpyxl; wb=openpyxl.load_workbook(f,read_only=True,data_only=True); parts=[]
            for nm in wb.sheetnames:
                rs=[[("" if v is None else v) for v in row] for row in wb[nm].iter_rows(values_only=True)]
                rs=[r for r in rs if any(str(x).strip() for x in r)]
                if rs: parts.append(f"### sheet: {nm}\n\n"+md_table(rs))
            wb.close(); body="\n\n".join(parts)
        elif ext==".pdf":
            import fitz; d=fitz.open(f); parts=[]; low=0
            md=os.path.join(DST,os.path.splitext(os.path.relpath(f,ROOT))[0]+".md")
            assets=os.path.join(DST,os.path.splitext(os.path.relpath(f,ROOT))[0]+".assets")
            for pi in range(d.page_count):
                t=d[pi].get_text().strip()
                if len(t)<40:
                    low+=1; os.makedirs(assets,exist_ok=True); png=os.path.join(assets,f"page_{pi+1:03d}.png")
                    try: d[pi].get_pixmap(dpi=150).save(png); parts.append(f"### page {pi+1}(少字,渲PNG)\n\n![p]({os.path.relpath(png,os.path.dirname(assets))})\n\n```\n{t}\n```")
                    except Exception: parts.append(f"### page {pi+1}\n\n```\n{t}\n```")
                else: parts.append(f"### page {pi+1}\n\n{clean(t)}")
            if low: loss.append(f"pdf_low_text={low}")
            d.close(); body="\n\n".join(parts)
        st="ok" if body and body.strip() else "empty"
    except Exception as e: err=repr(e)[:200]; loss.append("convert_error")
    return body,loss,st,err

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--types",required=True)
    ap.add_argument("--shard",required=True)  # i/N
    a=ap.parse_args()
    types={"."+t.strip().lower() for t in a.types.split(",")}
    i,N=[int(x) for x in a.shard.split("/")]
    os.makedirs(DST,exist_ok=True)
    files=[p for p in glob.glob(os.path.join(ROOT,"**","*"),recursive=True)
           if os.path.isfile(p) and not os.path.basename(p).startswith("~$")
           and os.path.splitext(p)[1].lower() in types]
    files.sort()
    mine=[f for k,f in enumerate(files) if k%N==i]
    tag=a.types.replace(",","_")+f"_{i}of{N}"
    print(f"[{tag}] 本片 {len(mine)}/{len(files)} 个 -> {DST}",flush=True)
    rows=[]
    for j,f in enumerate(mine,1):
        rel=os.path.relpath(f,ROOT); ext=os.path.splitext(f)[1].lower()
        md=os.path.join(DST,os.path.splitext(rel)[0]+".md")
        body,loss,st,err=convert(f,ext)
        if st=="ok":
            os.makedirs(os.path.dirname(md),exist_ok=True)
            open(md,"w",encoding="utf-8").write(header(rel,f,ext.lstrip("."),st,loss)+body)
        rows.append({"rel":rel,"type":ext.lstrip("."),"status":st,"text_len":len(body),"loss_flags":";".join(loss),"md":md if st=="ok" else "","error":err})
        if j%10==0 or j==len(mine): print(f"  [{tag}] {j}/{len(mine)} {st} {rel[-50:]}",flush=True)
    if _word:
        try: _word.Quit()
        except Exception: pass
    man=os.path.join(DST,f"_manifest_{tag}.csv")
    with open(man,"w",encoding="utf-8-sig",newline="") as fh:
        w=csv.DictWriter(fh,fieldnames=["rel","type","status","text_len","loss_flags","md","error"]); w.writeheader(); w.writerows(rows)
    from collections import Counter
    print(f"[{tag}] 完成 {dict(Counter(r['status'] for r in rows))} -> {man}",flush=True)

if __name__=="__main__":
    main()
